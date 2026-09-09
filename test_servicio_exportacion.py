"""Regresiones del servicio reutilizable de exportación."""
import csv
from io import BytesIO, TextIOWrapper
import zipfile

import openpyxl
import pandas as pd
import pytest

import base_datos
import consultas_boe
import servicio_exportacion as exportacion
from test_consultas_boe import ruta_busqueda


@pytest.mark.parametrize("filtros", [
    {},
    {"texto": "ingeniero"},
    {"fecha_desde": "2025-02-01", "fecha_hasta": "2025-02-01"},
    {"comunidad_autonoma": "Andalucía"},
    {"provincia": "Madrid"},
    {"municipio": "Madr"},
    {"municipio_exacto": "Madrid", "municipio_provincia_exacto": "Madrid"},
    {"administracion": "Ayuntamiento A"},
    {"ambito": "LOCAL"},
    {"tipo_entidad": "MUNICIPAL"},
    {"sistema": "Concurso"},
    {"turno": "Libre"},
    {"escala": "E1"},
    {"subescala": "S1"},
    {"clase": "C1"},
])
def test_exportacion_filtrada_reutiliza_el_universo_del_listado(ruta_busqueda, filtros):
    assert len(exportacion.obtener_oposiciones_filtradas(ruta_busqueda, **filtros)) == (
        consultas_boe.buscar_oposiciones(ruta_busqueda, **filtros)["total"]
    )


@pytest.mark.parametrize("orden", sorted(consultas_boe._ORDEN_BUSQUEDA))
def test_exportacion_filtrada_respeta_los_ordenes_permitidos(ruta_busqueda, orden):
    exportadas = exportacion.obtener_oposiciones_filtradas(ruta_busqueda, orden=orden)
    listado = consultas_boe.buscar_oposiciones(
        ruta_busqueda, orden=orden, tamano_pagina=100
    )
    assert exportadas["Puesto"].tolist() == [fila["puesto"] for fila in listado["filas"]]


def test_exportacion_filtrada_ignora_estado_visual_y_rechaza_orden_invalido(ruta_busqueda):
    normal = exportacion.obtener_oposiciones_filtradas(ruta_busqueda, orden="fecha_desc")
    visual = exportacion.obtener_oposiciones_filtradas(
        ruta_busqueda,
        orden="fecha_desc",
        pagina=9,
        tamano_pagina=1,
        ver_todas=True,
        vista="mapa",
        sin_coordenadas=True,
        pagina_sin_coordenadas=3,
    )
    pd.testing.assert_frame_equal(normal, visual)
    with pytest.raises(ValueError, match="Orden no permitido"):
        exportacion.obtener_oposiciones_filtradas(ruta_busqueda, orden="DROP TABLE")


def test_csv_zip_completo_tiene_cinco_datasets_y_protege_formulas(ruta_busqueda, tmp_path):
    conexion = base_datos.conectar(ruta_busqueda)
    try:
        conexion.execute("UPDATE oposiciones SET puesto = '=SUM(1;1)' WHERE oposicion_id = 1")
        conexion.commit()
    finally:
        conexion.close()
    salida = tmp_path / "base.zip"
    resultado = exportacion.exportar_base_csv_zip(ruta_busqueda, salida)
    assert resultado["datasets"]["Oposiciones"] == 3
    esperados = exportacion.preparar_datasets_para_xlsx(
        exportacion.cargar_datos_exportacion_completa(ruta_busqueda)[0]
    )
    with zipfile.ZipFile(salida) as archivo:
        assert archivo.namelist() == list(exportacion.NOMBRES_CSV_COMPLETOS.values())
        datos = archivo.read("oposiciones.csv")
        assert datos.startswith(b"\xef\xbb\xbf")
        filas = list(csv.DictReader(TextIOWrapper(BytesIO(datos), encoding="utf-8-sig"), delimiter=";"))
        for nombre, archivo_csv in exportacion.NOMBRES_CSV_COMPLETOS.items():
            leido = pd.read_csv(
                BytesIO(archivo.read(archivo_csv)), sep=";", encoding="utf-8-sig"
            )
            assert leido.columns.tolist() == esperados[nombre].columns.tolist()
            assert len(leido) == len(esperados[nombre])
    assert filas[0]["Puesto"] == "'=SUM(1;1)"
    assert "Puesto_normalizado" in filas[0]
    conexion = base_datos.conectar(ruta_busqueda, readonly=True)
    try:
        assert conexion.execute("SELECT puesto FROM oposiciones WHERE oposicion_id = 1").fetchone()[0] == "=SUM(1;1)"
    finally:
        conexion.close()


def test_xlsx_y_csv_filtrados_comparten_dataset_formato_y_formula(ruta_busqueda, tmp_path):
    conexion = base_datos.conectar(ruta_busqueda)
    try:
        conexion.execute(
            "UPDATE oposiciones SET puesto = ?, administracion = ?, escala = ?, clase = ? WHERE oposicion_id = 1",
            ("=SUM(1;1)", "+cmd", "-1+1", "@algo"),
        )
        conexion.commit()
    finally:
        conexion.close()
    xlsx, csv_path = tmp_path / "filtradas.xlsx", tmp_path / "filtradas.csv"
    resultado_xlsx = exportacion.exportar_oposiciones_filtradas_xlsx(
        ruta_busqueda, xlsx, orden="fecha_asc"
    )
    resultado_csv = exportacion.exportar_oposiciones_filtradas_csv(
        ruta_busqueda, csv_path, orden="fecha_asc"
    )
    assert resultado_xlsx["columnas"] == exportacion.COLUMNAS_OPOSICIONES_FILTRADAS
    assert resultado_csv["columnas"] == exportacion.COLUMNAS_OPOSICIONES_FILTRADAS
    desde_xlsx = pd.read_excel(xlsx, sheet_name="Oposiciones", keep_default_na=False)
    desde_csv = pd.read_csv(csv_path, sep=";", encoding="utf-8-sig", keep_default_na=False)
    pd.testing.assert_frame_equal(desde_xlsx, desde_csv, check_dtype=False)
    assert desde_csv.loc[0, "Puesto"] == "'=SUM(1;1)"
    assert desde_csv.loc[0, "Administración"] == "'+cmd"
    assert desde_csv.loc[0, "Escala"] == "'-1+1"
    assert desde_csv.loc[0, "Clase"] == "'@algo"

    libro = openpyxl.load_workbook(xlsx)
    hoja = libro["Oposiciones"]
    assert libro.sheetnames == ["Oposiciones"]
    assert libro.active.title == "Oposiciones"
    assert hoja.freeze_panes == "A2" and hoja.auto_filter.ref == hoja.dimensions
    assert all(celda.font.bold for celda in hoja[1])
    assert hoja.cell(2, exportacion.COLUMNAS_OPOSICIONES_FILTRADAS.index("Enlace BOE") + 1).hyperlink


def test_exportacion_filtrada_sin_resultados_es_valida(ruta_busqueda, tmp_path):
    xlsx, csv_path = tmp_path / "vacia.xlsx", tmp_path / "vacia.csv"
    exportacion.exportar_oposiciones_filtradas_xlsx(ruta_busqueda, xlsx, texto="inexistente")
    exportacion.exportar_oposiciones_filtradas_csv(ruta_busqueda, csv_path, texto="inexistente")
    libro = openpyxl.load_workbook(xlsx)
    hoja = libro["Oposiciones"]
    assert hoja.max_row == 1
    assert [celda.value for celda in hoja[1]] == exportacion.COLUMNAS_OPOSICIONES_FILTRADAS
    vacio = pd.read_csv(csv_path, sep=";", encoding="utf-8-sig")
    assert vacio.empty and vacio.columns.tolist() == exportacion.COLUMNAS_OPOSICIONES_FILTRADAS
